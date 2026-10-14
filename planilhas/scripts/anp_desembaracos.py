#!/usr/bin/env python3
"""
Relatório de Desembaraços da ANP: quem importou cada carga de combustível.

Módulo compartilhado por oferta_diesel.py e oferta_gasolina.py. Não roda
sozinho. Além do leitor do relatório, guarda duas coisas que os dois scripts
precisam e que não cabiam em nenhum deles: o calendário de quem era dono de
cada refinaria em cada mês, e o perfil do fornecimento no formato da FGV.

O QUE É ESTE ARQUIVO NA ANP

A série volumétrica de comércio exterior diz quanto entrou no país, mas não
diz quem trouxe. Quem diz é este relatório, que a ANP monta do Siscomex e
publica num .xlsx por ano, com importador, CNPJ, NCM, porto de despacho, país
de origem e quantidade em quilos. **Começa em janeiro de 2020** — antes disso
não existe dado por agente.

Três armadilhas, todas tratadas aqui:

  - **O layout muda de ano para ano.** Em 2025 a ANP inseriu a coluna "UF DO
    CNPJ*", que desloca tudo o que vem depois. Por isso as colunas são
    localizadas pelo nome, nunca pela posição.
  - **Vem em quilos.** A conversão para m³ usa 0,84 t/m³ no diesel e 0,74 na
    gasolina, que não é escolha nossa: é o fator que a própria ANP aplica. A
    densidade implícita entre este arquivo e a série volumétrica dá exatamente
    esses valores em quase todos os meses.
  - **Alguns meses saem incompletos** e são completados depois. Quem usa tem
    que comparar com o volumétrico do mesmo mês — é o que a coluna "cobertura"
    das planilhas faz.

E a ressalva que nenhum tratamento resolve: o relatório nomeia o **importador
de registro**, não o dono final da carga. Trading que importa por conta de uma
distribuidora aparece como trading.
"""
import datetime
import os
import re
import unicodedata
import urllib.request

DESEMB = ("https://www.gov.br/anp/pt-br/assuntos/importacoes-e-exportacoes/"
          "arquivos-desembaracos/desembaraco-%d.xlsx")
INICIO = 2020
NAVEGADOR = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36"
                           " (KHTML, like Gecko) Chrome/131.0 Safari/537.36"}

NCM_DIESEL = "27101921"      # gasóleo (óleo diesel); é o único NCM de diesel
NCM_GASOLINA = "27101259"    # "outras gasolinas", que é a gasolina A
DENSIDADE = {NCM_DIESEL: 0.84, NCM_GASOLINA: 0.74}


def baixar(ano, cache):
    """O .xlsx de um ano, ou None se a ANP ainda não publicou."""
    os.makedirs(cache, exist_ok=True)
    alvo = os.path.join(cache, "desembaraco-%d.xlsx" % ano)
    if not os.path.exists(alvo) or os.path.getsize(alvo) < 10000:
        print("  baixando desembaraco-%d.xlsx…" % ano)
        try:
            req = urllib.request.Request(DESEMB % ano, headers=NAVEGADOR)
            with urllib.request.urlopen(req, timeout=300) as r:
                conteudo = r.read()
        except Exception as e:
            print("    sem arquivo para %d (%s)" % (ano, e))
            return None
        with open(alvo, "wb") as f:
            f.write(conteudo)
    return alvo


def linhas(ano, cache):
    """Cada desembaraço do ano, como dicionário. Colunas pelo nome, porque a
    posição muda de um ano para outro."""
    from openpyxl import load_workbook
    arq = baixar(ano, cache)
    if not arq:
        return
    wb = load_workbook(arq, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    col = None
    for linha in ws.iter_rows(values_only=True):
        if col is None:
            if linha and linha[0] == "Mês de desembaraço":
                nomes = [str(c).strip() if c else "" for c in linha]
                col = {
                    "imp": nomes.index("Importador"),
                    "cnpj": nomes.index("CNPJ"),
                    "ncm": nomes.index("NCM"),
                    "kg": next(i for i, n in enumerate(nomes)
                               if n.startswith("Quantidade")),
                    "uf": nomes.index("UF DO CNPJ*") if "UF DO CNPJ*" in nomes else None,
                }
            continue
        if not isinstance(linha[0], int):
            continue
        quem, kg = linha[col["imp"]], linha[col["kg"]]
        if not quem or not isinstance(kg, (int, float)):
            continue
        yield {
            "mes": "%d-%02d" % (ano, linha[0]),
            "importador": str(quem).strip().upper(),
            "cnpj": re.sub(r"\D", "", str(linha[col["cnpj"]] or "")),
            "ncm": str(linha[col["ncm"]]).strip(),
            "uf": (str(linha[col["uf"]]).strip() if col["uf"] is not None
                   and linha[col["uf"]] else None),
            "kg": kg,
        }
    wb.close()


def anos(cache):
    ano = INICIO
    while ano <= datetime.date.today().year:
        if not baixar(ano, cache):
            return
        yield ano
        ano += 1


def por_importador(ncm, cache):
    """{"AAAA-MM": {importador: m³}} do produto pedido."""
    dens = DENSIDADE[ncm]
    fora = {}
    for ano in anos(cache):
        for r in linhas(ano, cache):
            if r["ncm"] != ncm:
                continue
            fora.setdefault(r["mes"], {}).setdefault(r["importador"], 0.0)
            fora[r["mes"]][r["importador"]] += r["kg"] / 1000.0 / dens
    return fora


# ---------------------------------------------------------------- classificação

# A ANP só dá a razão social, então o grupo vem do nome mais o que se sabe da
# estrutura societária. Vários grupos atravessam categorias — a Vibra tem
# distribuidora e trading, a Atem tem refinaria e distribuidora, Royal FIC e
# Ciapetro têm os dois braços — e aqui o braço trading vai junto com o grupo
# dele, não com as tradings independentes. É escolha discutível: a aba
# "Importação por empresa" das planilhas mostra o dado cru, sem classificação.
GRUPOS = (
    ("Petrobras", ("PETROLEO BRASILEIRO",)),
    ("refinaria privada", ("REFINARIA DE MANAUS", "REFINARIA DE PETROLEOS DE MANGUINHOS",
                           "REFINARIA DE MATARIPE", "REFINARIA DE PETROLEO RIOGRANDENSE",
                           "3R POTIGUAR", "DAX OIL REFINO")),
    ("central petroquímica", ("BRASKEM",)),
    ("outros/uso próprio", ("AGGREKO", "ECORI ENERGIA SOLAR", "BSBIOS", "INPASA",
                            "TEMAPE", "ICE QUIMICA", "HSDC COMERCIO")),
    ("grupo distribuidor", ("VIBRA ", "IPIRANGA PRODUTOS", "ATEM'S DISTRIBUIDORA",
                            "ROYAL FIC", "CIAPETRO", "ALESAT", "DISLUB",
                            "SETTA COMBUSTIVEIS", "FEDERAL ENERGIA", "TDC DISTRIBUIDORA",
                            "OLIVEIRA ENERGIA", "NFX COMBUSTIVEIS", "POTENCIAL PETROLEO",
                            "DISTRIBUIDORA", "PETROTORQUE", "REJAILE", "RUMOS", "ON PETRO")),
)
# A ordem de GRUPOS é prioridade de busca; CLASSES é a ordem das colunas.
CLASSES = ["Petrobras", "grupo distribuidor", "trading/importador independente",
           "refinaria privada", "central petroquímica", "outros/uso próprio"]
assert set(CLASSES) == {n for n, _ in GRUPOS} | {"trading/importador independente"}


def classe(nome):
    for grupo, marcas in GRUPOS:
        if any(m in nome for m in marcas):
            return grupo
    return "trading/importador independente"


def nome_legivel(razao):
    """ATEM'S DISTRIBUIDORA → Atem's Distribuidora."""
    return razao.title().replace("'S", "'s")


# --------------------------------------------------------------------- regiões

def sem_acento(s):
    return "".join(c for c in unicodedata.normalize("NFD", s or "")
                   if unicodedata.category(c) != "Mn").upper().strip()


REGIAO_DA_UF = {}
for _sigla, _ufs in (
    ("NO", "ACRE AMAPA AMAZONAS PARA RONDONIA RORAIMA TOCANTINS"),
    ("NE", "ALAGOAS BAHIA CEARA MARANHAO PARAIBA PERNAMBUCO PIAUI SERGIPE"),
    ("SE", "ESPIRITO SANTO|MINAS GERAIS|RIO DE JANEIRO|SAO PAULO"),
    ("SU", "PARANA|RIO GRANDE DO SUL|SANTA CATARINA"),
    ("CO", "DISTRITO FEDERAL|GOIAS|MATO GROSSO|MATO GROSSO DO SUL"),
):
    for _uf in (_ufs.split("|") if "|" in _ufs else _ufs.split()):
        REGIAO_DA_UF[_uf.strip()] = _sigla
REGIAO_DA_UF["RIO GRANDE DO NORTE"] = "NE"
REGIAO_DA_UF["ESPIRITO SANTO"] = "SE"
REGIOES = ["NO", "NE", "SE", "SU", "CO"]


def regiao(uf):
    """Região de uma UF escrita como a ANP escreve, com ou sem acento.
    Devolve "ND" para o que não dá para localizar."""
    return REGIAO_DA_UF.get(sem_acento(uf), "ND")


def mapa_uf(cache):
    """{CNPJ: UF} e {raiz do CNPJ: UF}, montados dos anos que trazem a coluna
    "UF DO CNPJ" (2025 em diante) para retropolar os que não trazem.

    O artigo da FGV que inspirou o recorte regional aloca a importação direta à
    região da SEDE da empresa, não ao porto de entrada — e é isso que esta
    coluna dá. Cobertura da retropolação: ~73% do volume em 2020, 88% a 96%
    entre 2021 e 2024, 100% de 2025 em diante. O que sobra fica em "ND"."""
    import collections
    cheio, raiz = {}, collections.defaultdict(collections.Counter)
    for ano in anos(cache):
        for r in linhas(ano, cache):
            u = r["uf"]
            if not u or "NAO REGULADO" in sem_acento(u):
                continue
            cheio[r["cnpj"]] = u
            raiz[r["cnpj"][:8]][u] += r["kg"]
    return cheio, {k: v.most_common(1)[0][0] for k, v in raiz.items()}


def uf_do_importador(r, cheio, raiz):
    """A UF de um desembaraço: a declarada, ou a do mapa, ou None."""
    if r["uf"] and "NAO REGULADO" not in sem_acento(r["uf"]):
        return r["uf"]
    return cheio.get(r["cnpj"]) or raiz.get(r["cnpj"][:8])


# --------------------------------------------- de quem era a planta, e quando

# A ANP renomeia as plantas RETROATIVAMENTE: a de Mataripe aparece como REFMAT
# desde 2000 e a de Manaus como REAM, mesmo nos anos em que eram RLAM e REMAN
# da Petrobras. Por isso o dono não pode sair do nome — sai da data em que a
# venda fechou, pelos comunicados da própria Petrobras:
#
#   RLAM  -> Acelen / Refinaria de Mataripe : 30/11/2021
#   REMAN -> Atem / Refinaria de Manaus     : 30/11/2022
#   RPCC  -> 3R Potiguar (Polo Potiguar)    : ANP aprovou em 07/06/2023
#
NUNCA_PETROBRAS = {"MANGUINHOS", "RIOGRANDENSE", "DAX OIL", "UNIVEN", "SSOIL"}
VENDIDAS = {
    "REFMAT": (2021, 11),
    "REAM": (2022, 11),
    "3R POTIGUAR (EX-3R POTIGUAR (EX-RPCC))": (2023, 5),
}


def dono(planta, ano, mes):
    """"petrobras" ou "outras" para uma refinaria num mês."""
    n = planta.strip().upper()
    if n in NUNCA_PETROBRAS:
        return "outras"
    if n in VENDIDAS:
        return "petrobras" if (ano, mes) <= VENDIDAS[n] else "outras"
    return "petrobras"


# ------------------------------------------------------ perfil do fornecimento

# As seis linhas da tabela, na ordem em que a FGV as publica. A regra do
# asterisco dela ("inclui as importações realizadas pelas controladoras") é o
# que faz a importação da Petrobras subir para a linha "Petrobras" e a das
# refinarias privadas para "Outras", em vez de ficarem em "Importações".
LINHAS_PERFIL = ["Refinarias e Centrais Petroquímicas", "Petrobras", "Outras",
                 "Importações", "Tradings", "Distribuidores"]


def perfil(producao, importacao, meses, janela=12):
    """Perfil do fornecimento por região, em % do volume, no formato da tabela
    T.1 de "Impacto da defasagem da gasolina no IPCA e reflexos de sua
    subvenção nos resultados da Petrobras" (FGV).

      producao:   {(mes, "petrobras"|"outras", regiao): m³}
      importacao: {(mes, classe, regiao): m³}
      janela:     meses somados para cada ponto (12 = a janela do artigo, que
                  mata a sazonalidade; 1 = mês puro)

    Devolve uma linha por (mês final da janela, região), com as seis
    participações e o volume total do período.
    """
    import collections
    pos = {m: i for i, m in enumerate(meses)}
    fora = []
    for m in meses:
        i = pos[m]
        if i + 1 < janela:
            continue
        bloco = set(meses[i + 1 - janela:i + 1])
        acc = collections.defaultdict(lambda: collections.defaultdict(float))
        for (mm, g, reg), v in producao.items():
            if mm in bloco:
                acc[reg]["prod_" + g] += v
        for (mm, c, reg), v in importacao.items():
            if mm in bloco:
                acc[reg]["imp_" + c] += v
        # as regiões saem do próprio dado: na série nacional tudo vem marcado
        # como "Brasil" e só essa linha é emitida; no recorte regional vêm as
        # cinco regiões mais "ND" para o que não se consegue localizar.
        regs = sorted(r for r in acc if r != "Brasil")
        total = collections.defaultdict(float)
        for d in acc.values():
            for k, v in d.items():
                total[k] += v
        for rot, d in [("Brasil", total)] + [(r, acc[r]) for r in regs]:
            pb = d["prod_petrobras"] + d["imp_Petrobras"]
            ou = (d["prod_outras"] + d["imp_refinaria privada"]
                  + d["imp_central petroquímica"])
            tr = d["imp_trading/importador independente"]
            di = d["imp_grupo distribuidor"] + d["imp_outros/uso próprio"]
            tot = pb + ou + tr + di
            if not tot:
                continue
            fora.append([m, rot,
                         100.0 * (pb + ou) / tot, 100.0 * pb / tot, 100.0 * ou / tot,
                         100.0 * (tr + di) / tot, 100.0 * tr / tot, 100.0 * di / tot,
                         tot])
    return fora
