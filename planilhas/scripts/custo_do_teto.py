#!/usr/bin/env python3
"""
Monta custo-do-teto.xlsx — quanto custa à Petrobras vender abaixo da paridade
de importação, separando o que é prejuízo de caixa do que é receita não
realizada.

Uso:
    ~/.venvs/dados-economicos/bin/python scripts/custo_do_teto.py

Reaproveita defasagem.py (PPI da ANP, tabela da Petrobras, casamento semana a
semana) e as planilhas de oferta. NÃO entra na Action.

POR QUE SEPARAR AS DUAS CONTAS

Vender abaixo do PPI é um preço-teto de fato, e um teto cobra de duas formas
que não se parecem em nada:

  - **Custo de oportunidade**: sobre cada litro que a Petrobras produz e vende
    no país, ela deixa de arrecadar a diferença até a paridade. É receita que
    não entra, não dinheiro que sai. Uma estatal absorve isso por anos.
  - **Prejuízo de caixa**: quando ela própria importa para cobrir o mercado,
    compra perto da paridade e vende pelo preço de tabela. Aí a diferença sai
    do bolso, litro a litro.

É a diferença entre as duas áreas do diagrama de oferta e demanda com teto: a
faixa sobre a produção doméstica e a faixa sobre o volume importado. A primeira
é grande e suportável; a segunda é pequena e dói. Qual das duas aparece diz
quem está bancando o teto — e, portanto, quanto tempo ele dura.

O TESTE QUE ISSO PERMITE

Se o teto fosse insustentável, veríamos a Petrobras importando no vermelho. O
recorte por agente do desembaraço responde direto: nos meses de defasagem
aberta de 2026 ela importou ~zero de gasolina (distribuidores e tradings
trouxeram) e ~zero de diesel até junho, voltando a importar diesel em julho e
agosto, com a subvenção já em vigor. Ou seja: na gasolina o teto é pago em
receita não realizada; no diesel, em caixa, e desde julho com o Tesouro
dividindo a conta.

O QUE O NÚMERO NÃO É

  - O custo de oportunidade é **teto, não estimativa**: supõe que ela venderia
    o mesmo volume ao PPI. Não venderia — a demanda cairia, que é justamente o
    que o diagrama mostra. Serve para ordem de grandeza e comparação entre
    produtos, não como lucro perdido.
  - A defasagem mensal é **média simples dos pontos de entrega**, não ponderada
    por volume: a ANP não publica volume por ponto. Nos nove pontos a dispersão
    é pequena (frete), então a média simples não distorce.
  - A subvenção é paga a quem vende abaixo da paridade, e por isso abate o
    prejuízo de quem importa. Quanto exatamente chegou à Petrobras por litro
    não é público; a coluna líquida supõe o abatimento integral e é, portanto,
    o piso da exposição dela.
"""
import collections
import datetime
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import defasagem as df                                     # noqa: E402
import precos_combustiveis as pc                            # noqa: E402

SAIDA = os.path.join(pc.RAIZ, "custo-do-teto.xlsx")

# produto na tabela da Petrobras, aba do PPI da ANP, padrão do PDF, planilha de
# oferta e rótulo. O diesel da ANP é S10 ("Diesel A S10", no alto da aba), e é
# com o S10 da Petrobras que ele se compara.
PRODUTOS = (
    ("Gasolina A", "Gasolina R$ semanal", "Gasolina", "oferta-gasolina.xlsx",
     "gasolina"),
    ("Diesel S-10", "Diesel R$ semanal", "Diesel", "oferta-diesel.xlsx",
     "diesel"),
)

INICIO = "2026-01"          # o ano corrente inteiro, para o sinal virar à vista
CHOQUE = "2026-03"          # o mês em que a defasagem abriu, e a janela do artigo


def defasagem_mensal(produto, aba_ppi, padrao_pdf):
    """{mês: (defasagem R$/l, subvenção R$/l, nº de pontos, nº de semanas)}.

    Defasagem com sinal invertido em relação a defasagem.py: aqui positivo é
    quanto falta para a paridade, que é o que se multiplica por volume."""
    ppi = df.ler_ppi(df.baixar_ppi(), aba_ppi)
    caminho = pc.pdf_de(padrao_pdf)
    if not caminho:
        raise SystemExit("falta a tabela de %s da Petrobras" % produto)
    linhas = pc.ler_petrobras(caminho, produto=produto)
    escadas = {p: df.escadas_de(linhas, "Petrobras", p, produto=produto)
               for p in df.PONTOS}
    escadas = {p: e for p, e in escadas.items() if e[0] or e[1]}
    if not escadas:
        raise SystemExit("nenhum ponto de %s casou com PONTOS" % produto)
    ultima = max(d for tab, sub in escadas.values() for d, _v in tab + sub)

    por_mes = collections.defaultdict(list)
    for ponto, fim, tabela, subv, p_ppi, dif, _pc1, _pc2, _ext in \
            df.montar(ppi, escadas, ultima, {p: df.PONTOS[p] for p in escadas}):
        por_mes["%d-%02d" % (fim.year, fim.month)].append((ponto, -dif, subv))

    fora = {}
    for mes, itens in por_mes.items():
        gaps = [g for _p, g, _s in itens]
        # A subvenção é valor por litro fixado em lei, igual em todo ponto — a
        # média entra só sobre os pontos que PUBLICAM a coluna com desconto.
        # Médias sobre os nove pontos dariam R$ 0,34/l na gasolina em vez dos
        # R$ 0,44 da MP 1.358, diluídos por Guamaré e Manaus, onde a Petrobras
        # não republicou a tabela subvencionada.
        subs = [s for _p, _g, s in itens if s > 0]
        fora[mes] = (sum(gaps) / len(gaps),
                     (sum(subs) / len(subs)) if subs else 0.0,
                     len({p for p, _g, _s in itens}), len(itens))
    return fora, ultima, sorted(escadas)


def oferta(arquivo):
    """({mês: (% Petrobras, volume total m³)}, {mês: importação própria m³})."""
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(pc.RAIZ, arquivo), data_only=True)

    ws = wb["Perfil mensal"]
    grade = [list(r) for r in ws.iter_rows(values_only=True)]
    cab = [c for c in grade[0]]
    i_pb, i_vol = cab.index("Petrobras"), len(cab) - 1
    perfil = {str(r[0])[:7]: (r[i_pb], r[i_vol]) for r in grade[1:] if r[0]}

    ws = wb["Importação por agente"]
    grade = [list(r) for r in ws.iter_rows(values_only=True)]
    i_pb = grade[0].index("Petrobras")
    imp = {str(r[0])[:7]: (r[i_pb] or 0) for r in grade[1:] if r[0]}
    return perfil, imp


def montar_produto(produto, aba_ppi, padrao_pdf, arquivo_oferta, rotulo):
    gaps, ultima, pontos = defasagem_mensal(produto, aba_ppi, padrao_pdf)
    perfil, imp = oferta(arquivo_oferta)

    fora = []
    for mes in sorted(m for m in gaps if m >= INICIO and m in perfil):
        gap, subv, n_pontos, n_semanas = gaps[mes]
        pct_pb, volume = perfil[mes]
        if pct_pb is None or volume is None:
            continue
        vol_pb = volume * pct_pb / 100
        propria = imp.get(mes, 0)
        # m³ × R$/l = R$ mil; /1000 = R$ milhões
        oportunidade = gap * vol_pb / 1000
        caixa = gap * propria / 1000
        caixa_liq = max(gap - subv, 0) * propria / 1000
        fora.append([mes, gap, subv, vol_pb, propria, oportunidade, caixa,
                     caixa_liq, n_pontos, n_semanas])
    return fora, ultima, pontos


def escrever(resultados, hoje):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)

    leia = wb.create_sheet("LEIA-ME")
    for linha in [
        ["O custo de vender abaixo da paridade de importação"],
        ["Gerado em", hoje.strftime("%d/%m/%Y")],
        [],
        ["Aba", "O que é"],
        ["Custo do teto — gasolina", "Mês a mês, desde janeiro de 2026"],
        ["Custo do teto — diesel", "O mesmo, com o diesel A S10"],
        ["Resumo", "A soma da janela nos dois produtos, lado a lado"],
        [],
        ["UM TETO COBRA DE DUAS FORMAS, E ELAS NÃO SE PARECEM"],
        [],
        ["Custo de oportunidade", "Defasagem × volume que a Petrobras fornece no "
                                  "país. Receita que não entra."],
        ["", "Uma estatal absorve isso por anos: não há saída de caixa, e o "
             "resultado contábil só mostra o lucro menor."],
        ["Prejuízo de caixa", "Defasagem × o que a Petrobras ELA MESMA importou. "
                              "Compra perto da paridade e vende pela tabela."],
        ["", "Aí o dinheiro sai. É a conta que aperta, e é pequena enquanto ela "
             "não precisa importar."],
        [],
        ["São as duas áreas do diagrama de oferta e demanda com preço-teto: a "
         "faixa sobre a produção doméstica e a"],
        ["faixa sobre o volume importado. Qual das duas aparece diz quem está "
         "bancando o teto, e quanto tempo ele dura."],
        [],
        ["O CUSTO DE OPORTUNIDADE É TETO, NÃO ESTIMATIVA. Supõe que ela venderia o "
         "mesmo volume ao PPI, e não venderia:"],
        ["a demanda cairia, que é justamente o que o diagrama mostra. Serve para "
         "ordem de grandeza e para comparar os dois"],
        ["produtos — não como lucro perdido."],
        [],
        ["DEFASAGEM = PPI − PREÇO DE TABELA DA PETROBRAS, em R$/litro, média "
         "simples dos pontos de entrega em que as duas"],
        ["pontas existem. Sinal invertido em relação ao defasagem-gasolina.xlsx: "
         "aqui positivo é quanto FALTA para a paridade,"],
        ["que é o que faz sentido multiplicar por volume. A ANP não publica volume "
         "por ponto, então a média é simples; a"],
        ["dispersão entre os pontos é de frete e não distorce."],
        [],
        ["A SUBVENÇÃO é paga a quem vende abaixo da paridade, e por isso abate o "
         "prejuízo de quem importa. Quanto chegou à"],
        ["Petrobras por litro não é público. A coluna 'caixa, líquido da subvenção' "
         "supõe abatimento integral, e é portanto o"],
        ["PISO da exposição dela; a coluna bruta é o teto."],
        [],
        ["VOLUME FORNECIDO vem da aba 'Perfil mensal' das planilhas de oferta (a "
         "fatia da Petrobras no fornecimento ×"],
        ["volume do período). IMPORTAÇÃO PRÓPRIA vem da aba 'Importação por "
         "agente', coluna Petrobras, que é o"],
        ["Relatório de Desembaraços da ANP, por agente declarante no Siscomex."],
        [],
        ["Fontes: ANP (PPI, produção, vendas, desembaraços) e Petrobras (tabela de "
         "preços, Resolução ANP 795/2019)."],
    ]:
        leia.append(linha)
    leia["A1"].font = Font(bold=True, size=13)
    for c in ("A4", "A9", "A10", "A12"):
        leia[c].font = Font(bold=True)
    leia.column_dimensions["A"].width = 34
    leia.column_dimensions["B"].width = 96

    cab = ["mês", "defasagem R$/l", "subvenção R$/l", "volume Petrobras (m³)",
           "importação própria (m³)", "custo de oportunidade (R$ mi)",
           "caixa, bruto (R$ mi)", "caixa, líquido da subvenção (R$ mi)",
           "pontos", "semanas × ponto"]
    resumo = [["produto", "janela", "custo de oportunidade (R$ mi)",
               "caixa, bruto (R$ mi)", "caixa, líquido (R$ mi)",
               "razão oportunidade / caixa bruto"]]

    def somar(linhas, desde=None):
        sel = [l for l in linhas if desde is None or l[0] >= desde]
        if not sel:
            return None
        o, cx, cl = (sum(l[i] for l in sel) for i in (5, 6, 7))
        return ["%s a %s" % (sel[0][0], sel[-1][0]), o, cx, cl,
                (o / cx) if cx else None]

    for rotulo, (linhas, ultima, pontos) in resultados.items():
        ws = wb.create_sheet("Custo do teto — " + rotulo)
        ws.append(cab)
        for l in linhas:
            ws.append(l)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        larguras = (9, 15, 15, 22, 24, 27, 20, 32, 8, 16)
        for j, w in enumerate(larguras, 1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = w
        for linha_ws in ws.iter_rows(min_row=2):
            for j, c in enumerate(linha_ws):
                if j in (1, 2):
                    c.number_format = "0.000"
                elif j in (3, 4):
                    c.number_format = "#,##0"
                elif j in (5, 6, 7):
                    c.number_format = "#,##0"
        ws.append([])
        ws.append(["Pontos de entrega: " + ", ".join(pontos)])
        ws.append(["Tabela mais recente da Petrobras: "
                   + ultima.strftime("%d/%m/%Y")])

        for janela, desde in (("", None), (" (desde o choque)", CHOQUE)):
            soma = somar(linhas, desde)
            if soma:
                resumo.append([rotulo + janela] + soma)

    ws = wb.create_sheet("Resumo")
    for l in resumo:
        ws.append(l)
    for c in ws[1]:
        c.font = Font(bold=True)
    for linha_ws in ws.iter_rows(min_row=2):
        for j, c in enumerate(linha_ws):
            if j in (2, 3, 4):
                c.number_format = "#,##0"
            elif j == 5:
                c.number_format = "0.0"
    for j, w in enumerate((12, 20, 27, 20, 22, 30), 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = w

    wb.save(SAIDA)


def main():
    hoje = datetime.date.today()
    resultados = {}
    for produto, aba, padrao, arquivo, rotulo in PRODUTOS:
        print("%s…" % produto)
        linhas, ultima, pontos = montar_produto(produto, aba, padrao, arquivo,
                                                rotulo)
        print("  %d pontos, %d meses de %s a %s"
              % (len(pontos), len(linhas), linhas[0][0], linhas[-1][0]))
        for l in linhas:
            print("    %s defasagem %5.3f (subv %5.3f) | fornece %10s m³, "
                  "importa %9s m³ | oportunidade %7s mi, caixa %6s mi"
                  % (l[0], l[1], l[2], format(round(l[3]), ","),
                     format(round(l[4]), ","), format(round(l[5]), ","),
                     format(round(l[6]), ",")))
        o, cx, cl = (sum(l[i] for l in linhas) for i in (5, 6, 7))
        print("  soma: oportunidade R$ %s mi, caixa bruto R$ %s mi, "
              "líquido R$ %s mi — razão %.0f×"
              % (format(round(o), ","), format(round(cx), ","),
                 format(round(cl), ","), o / cx if cx else float("nan")))
        resultados[rotulo] = (linhas, ultima, pontos)

    escrever(resultados, hoje)
    print("\n%s  (%.0f KB)"
          % (os.path.relpath(SAIDA, pc.RAIZ), os.path.getsize(SAIDA) / 1024))


if __name__ == "__main__":
    main()
